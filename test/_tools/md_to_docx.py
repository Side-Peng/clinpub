"""把 manuscript.md 转成 DOCX（仅用于查看效果，非交付格式）。"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

BASE = Path(r"D:\Projects\Skills_Makeing\02_Gsd Like Scientific pipeline\clinpub\test")
src = BASE / "05_Manuscript" / "manuscript.md"
out = BASE / "05_Manuscript" / "manuscript_preview.docx"


def set_font(run, east_asian: str = "宋体", size: float = 10.5, bold: bool = False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asian)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


text = src.read_text(encoding="utf-8")
text = re.sub(r"\A---.*?\n---\n", "", text, flags=re.S)
text = re.sub(r"<!--.*?-->", "", text, flags=re.S)

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)

for line in text.split("\n"):
    s = line.rstrip()
    if not s.strip():
        continue
    if s.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(s[2:].strip()), "黑体", 15, True)
    elif s.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        set_font(p.add_run(s[3:].strip()), "黑体", 12, True)
    elif s.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        set_font(p.add_run(s[4:].strip()), "黑体", 10.5, True)
    else:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.space_after = Pt(6)
        for part in re.split(r"(\*\*[^*]+\*\*)", s):
            if not part:
                continue
            if part.startswith("**"):
                set_font(p.add_run(part[2:-2]), "宋体", 10.5, True)
            else:
                set_font(p.add_run(part), "宋体", 10.5, False)

doc.save(out)
print(f"  已生成 {out}")
print(f"  段落数 {len(doc.paragraphs)}")
